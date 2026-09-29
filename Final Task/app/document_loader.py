import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Dict, List, Optional, Union

from pypdf import PdfReader
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

logger = logging.getLogger(__name__)


@dataclass
class Document:
    """Represents an ingested document."""
    content: str
    source: str
    metadata: Dict = field(default_factory=dict)


class DocumentLoader:
    """
    Universal document loader supporting PDF, TXT, and DOCX files.
    Can ingest from local file paths or in-memory file objects (Streamlit uploads).
    """

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}

    @classmethod
    def load_from_bytes(cls, file_bytes: bytes, filename: str) -> Document:
        """Load document from raw bytes and filename."""
        suffix = Path(filename).suffix.lower()

        if suffix == ".txt":
            content = cls._decode_text_bytes(file_bytes, filename)
        elif suffix == ".pdf":
            content = cls._load_pdf_from_stream(io.BytesIO(file_bytes), filename)
        elif suffix == ".docx":
            content = cls._load_docx_from_stream(io.BytesIO(file_bytes), filename)
        else:
            raise ValueError(
                f"Unsupported file format '{suffix}'. Supported formats: {', '.join(cls.SUPPORTED_EXTENSIONS)}"
            )

        cleaned_text = content.strip()
        if not cleaned_text:
            raise ValueError(f"The file '{filename}' does not contain readable text or is empty.")

        return Document(
            content=cleaned_text,
            source=filename,
            metadata={"source": filename, "format": suffix, "char_count": len(cleaned_text)},
        )

    @classmethod
    def load_from_path(cls, file_path: Union[str, Path]) -> Document:
        """Load document from a file path."""
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        file_bytes = path.read_bytes()
        doc = cls.load_from_bytes(file_bytes, path.name)
        doc.metadata["path"] = str(path)
        return doc

    @classmethod
    def _decode_text_bytes(cls, data: bytes, filename: str) -> str:
        """Try decoding text bytes with common encodings."""
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode text file '{filename}'.")

    @classmethod
    def _load_pdf_from_stream(cls, stream: BinaryIO, filename: str) -> str:
        """Extract text from PDF pages."""
        try:
            reader = PdfReader(stream)
            pages_text: List[str] = []
            for idx, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages_text.append(f"[Page {idx}]\n{page_text.strip()}")

            if not pages_text:
                raise ValueError("PDF has no readable text layers.")
            return "\n\n".join(pages_text)
        except Exception as exc:
            raise ValueError(f"Error reading PDF '{filename}': {exc}") from exc

    @classmethod
    def _load_docx_from_stream(cls, stream: BinaryIO, filename: str) -> str:
        """Extract text from DOCX stream."""
        if not HAS_DOCX:
            raise RuntimeError("python-docx is not installed.")
        try:
            doc = docx.Document(stream)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n\n".join(paragraphs)
        except Exception as exc:
            raise ValueError(f"Error reading DOCX '{filename}': {exc}") from exc

    @classmethod
    def load_directory(cls, dir_path: Union[str, Path]) -> List[Document]:
        """Load all supported documents from a directory."""
        path = Path(dir_path)
        if not path.is_dir():
            return []

        docs: List[Document] = []
        for file in sorted(path.iterdir()):
            if file.suffix.lower() in cls.SUPPORTED_EXTENSIONS:
                try:
                    docs.append(cls.load_from_path(file))
                except Exception as exc:
                    logger.warning("Skipping file %s: %s", file.name, exc)
        return docs
